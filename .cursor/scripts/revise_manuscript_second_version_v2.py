#!/usr/bin/env python3
"""Add interaction + diversity Results sections to manuscript (v2).

Reads:  Manuscript/btp_manuscript_second_version.docx
Writes: Manuscript/btp_manuscript_second_version_v2.docx
New / changed text is blue (RGB 0, 112, 192).
"""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "Manuscript" / "btp_manuscript_second_version.docx"
OUT = ROOT / "Manuscript" / "btp_manuscript_second_version_v2.docx"

BLUE = RGBColor(0, 112, 192)

INTERACTION_HEADING = "Interaction-associated shifts in the possum gut microbiome"
INTERACTION_BODY = (
    "The fitted devil × temperature interaction term identified a further layer of genome-specific response: "
    "119 genomes (21%) showed a positive interaction coefficient, 102 (18%) a negative interaction and 355 (62%) "
    "were neutral (Fig. 3). Interaction-positive genomes were again concentrated in Bacillota_A (n = 77), with "
    "notable representation of Lachnospiraceae (n = 23) and CAG-508 (n = 19), indicating that combined predator "
    "and thermal context was associated with reweighting within the Firmicutes core rather than wholesale community "
    "turnover. Functional volcano contrasts identified 16 FDR-significant GIFT elements between interaction-positive "
    "and interaction-negative genomes (12 enriched among interaction-positive, four among interaction-negative). "
    "Interaction-positive genomes were enriched in nitrogen-compound degradation (D06; e.g. D0608, D0601, D0607), "
    "amino-acid degradation (D0512), vitamin biosynthesis (B0711) and appendage traits (S0202), whereas "
    "interaction-negative genomes carried higher capacity for lipid and sugar degradation (D0102, D0302). Because "
    "devil density and temperature co-vary geographically within Tasmania, these interaction patterns should be "
    "interpreted as associational and partially confounded with the main-effect gradients."
)

DIVERSITY_HEADING = "Diversity-associated shifts in the possum gut microbiome"
DIVERSITY_BODY = (
    "Landscape diversity (land-cover heterogeneity around sampling sites) showed a similar magnitude of genome-level "
    "response to devil density and temperature: 133 genomes (23%) were positively associated, 109 (19%) negatively "
    "associated and 334 (58%) neutral. Diversity-positive genomes were dominated by Bacillota_A (n = 108), "
    "including CAG-508 (n = 36), Lachnospiraceae (n = 32) and Borkfalkiaceae (n = 24). Functional contrasts "
    "identified 69 FDR-significant GIFT elements (Fig. 3), but with a strongly asymmetric distribution: only three "
    "elements were enriched among diversity-positive genomes—spore formation (S0301), acetate biosynthesis (B0401) "
    "and lipoteichoic acid structure (S0104)—whereas 66 elements showed higher capacity in diversity-negative genomes, "
    "spanning amino-acid and aromatic biosynthesis (B02, B08), vitamin biosynthesis (B07) and multiple degradation "
    "pathways including antibiotic (D0901) and amino-acid degradation (D0501). This pattern suggests that higher "
    "landscape diversity was associated with genome subsets carrying somewhat leaner biosynthetic and degradative "
    "pathway completeness, alongside modest enrichment of spore and surface-structure traits, rather than a broad "
    "expansion of metabolic repertoire at genome resolution."
)

SEC343_HEADING = "3.4.3 | Interaction-associated shifts in the microbiome"
SEC343_BODY = (
    "The interaction term captured genomes whose responses to temperature depended on devil context "
    "(119 positive, 102 negative). Sixteen GIFT elements differed significantly between interaction-positive "
    "and interaction-negative genomes, enriched in nitrogen degradation and appendage traits among "
    "interaction-positive MAGs (Fig. 3). Interpretation remains cautious because devil density and temperature "
    "are spatially structured across Tasmania."
)

SEC344_HEADING = "3.4.4 | Diversity-associated shifts in the microbiome"
SEC344_BODY = (
    "Landscape diversity showed parallel genome-level filtering (133 positive, 109 negative). "
    "Volcano contrasts revealed 69 FDR-significant elements, dominated by higher pathway completeness in "
    "diversity-negative genomes (66 elements) and only three elements enriched among diversity-positive "
    "genomes (S0301, B0401, S0104). Diversity-positive MAGs were concentrated in CAG-508 and Lachnospiraceae, "
    "consistent with habitat-heterogeneity-associated reweighting within Bacillota_A."
)


def replace_whole_paragraph(paragraph, text: str, color: RGBColor = BLUE) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.color.rgb = color


def insert_paragraph_before(paragraph, text: str, color: RGBColor = BLUE):
    parent = paragraph._element.getparent()
    new_p = deepcopy(paragraph._element)
    parent.insert(parent.index(paragraph._element), new_p)
    from docx.text.paragraph import Paragraph

    new_para = Paragraph(new_p, paragraph._parent)
    replace_whole_paragraph(new_para, text, color=color)
    return new_para


def find_paragraph(doc, predicate):
    for para in doc.paragraphs:
        if predicate(para.text.strip()):
            return para
    return None


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Missing source manuscript: {SRC}")

    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))

    # --- §3.3 title ---
    p_33 = find_paragraph(doc, lambda t: t.startswith("3.3 |"))
    if p_33:
        replace_whole_paragraph(
            p_33,
            "3.3 | Genome-resolved associations with devil density, temperature, interaction and landscape diversity",
        )

    # --- §3.3 new subsections before Figure 3 ---
    fig3 = find_paragraph(
        doc,
        lambda t: t.startswith("Figure 3. HMSC functional associations across four predictors"),
    )
    if fig3 is None:
        raise SystemExit("Could not find Figure 3 caption.")

    for text in (INTERACTION_HEADING, INTERACTION_BODY, DIVERSITY_HEADING, DIVERSITY_BODY):
        insert_paragraph_before(fig3, text)

    # --- §3.4 replace stale temperature duplicate with interaction + diversity summaries ---
    stale = find_paragraph(
        doc,
        lambda t: t.startswith("Functional contrasts mirrored this weaker signal"),
    )
    if stale is not None:
        p_after = None
        for i, para in enumerate(doc.paragraphs):
            if para._element is stale._element and i + 1 < len(doc.paragraphs):
                p_after = doc.paragraphs[i + 1]
                break
        if p_after is None:
            p_after = stale
        replace_whole_paragraph(stale, SEC343_HEADING)
        for text in (SEC343_BODY, SEC344_HEADING, SEC344_BODY):
            insert_paragraph_before(p_after, text)

    doc.save(str(OUT))
    print(f"Saved: {OUT}")
    print("New sections and revised text are coloured blue (RGB 0, 112, 192).")


if __name__ == "__main__":
    main()
