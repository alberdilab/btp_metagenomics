#!/usr/bin/env python3
"""Align btp_manuscript_second_version.docx with current pipeline outputs.

New / revised wording is rendered in blue (RGB 0, 112, 192) for easy review.
"""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "Manuscript" / "btp_manuscript_second_version.docx"
BACKUP = ROOT / "Manuscript" / "btp_manuscript_second_version_pre_revision.docx"
OUT = SRC

BLUE = RGBColor(0, 112, 192)


def replace_whole_paragraph(paragraph, text: str, color: RGBColor = BLUE) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.color.rgb = color


def replace_substring(paragraph, old: str, new: str, color: RGBColor = BLUE) -> bool:
    text = paragraph.text
    if old not in text:
        return False
    before, _, after = text.partition(old)
    paragraph.clear()
    if before:
        paragraph.add_run(before)
    run_new = paragraph.add_run(new)
    run_new.font.color.rgb = color
    if after:
        paragraph.add_run(after)
    return True


def insert_paragraph_before(paragraph, text: str, color: RGBColor = BLUE):
  parent = paragraph._element.getparent()
  new_p = deepcopy(paragraph._element)
  parent.insert(parent.index(paragraph._element), new_p)
  from docx.text.paragraph import Paragraph
  new_para = Paragraph(new_p, paragraph._parent)
  replace_whole_paragraph(new_para, text, color=color)
  return new_para


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Missing manuscript: {SRC}")

    shutil.copy2(SRC, BACKUP)
    doc = Document(str(SRC))
    p = doc.paragraphs

    # --- Abstract ---
    replace_substring(
        p[1],
        "Additionally, hierarchical modelling of species communities (HMSC) confirms temperature and devil density as important predictors.",
        "Hierarchical modelling of species communities (HMSC) further identified genome-level associations with devil density, temperature, landscape diversity, and their interaction across Tasmania.",
    )

    # --- Introduction ---
    replace_substring(
        p[9],
        "To address this, we analysed faecal samples of 72 wild brushtail possums collected across Tasmania using metagenomic assembly and binning.",
        "To address this, we collected faecal material from 73 wild brushtail possums across Tasmania and mainland Australia (72 retained after quality filtering; sample EHI01340 excluded) using metagenomic assembly and binning.",
    )

    # --- Methods: HMSC ---
    replace_whole_paragraph(
        p[44],
        "Hierarchical modelling of species communities (HMSC) was applied genome-by-genome to 576 MAGs detected in 55 Tasmania samples. "
        "Fixed effects comprised devil density, mean annual temperature, landscape diversity (land-cover heterogeneity), "
        "log-transformed sequencing depth, and a devil × temperature interaction. Random effects comprised sampling region "
        "and individual animal. Genome-level association classes (positive, negative, neutral) were assigned from posterior "
        "support (threshold 0.9). Functional responses of GIFT elements were contrasted between positive and negative genomes "
        "and summarised in four-panel volcano plots (devil density, temperature, interaction, landscape diversity).",
    )

    # --- Results 3.1 ---
    replace_substring(
        p[49],
        "The assembly and binning of 352.3 Gb",
        "From 73 collected faecal samples (72 analysed), the assembly and binning of 352.3 Gb",
    )

    # --- Results 3.2: prepend alpha-diversity paragraph ---
    alpha_para = (
        "Across all 72 analysed samples, possum gut microbiomes harboured on average 307 ± 138 MAGs (richness), "
        "with abundance-weighted diversity of 115 ± 69, phylogenetic diversity of 7.4 ± 1.4, and functional diversity "
        "of 1.44 ± 0.05. Linear mixed-effects models with broad environment as a fixed effect and region as a random "
        "intercept indicated significant environment effects on richness (F4,52 = 2.74, p = 0.038) and phylogenetic "
        "diversity (F4,52 = 3.77, p = 0.009), but not on abundance-weighted diversity (F4,52 = 1.48, p = 0.221) or "
        "functional diversity (F4,52 = 0.38, p = 0.821; Fig. 2)."
    )
    insert_paragraph_before(p[57], alpha_para)

    # Re-fetch paragraphs after insertion
    p = doc.paragraphs
    for i, para in enumerate(p):
        if para.text.strip().startswith("PERMANOVA showed that broad environment"):
            replace_whole_paragraph(
                para,
                "PERMANOVA showed that broad environment significantly structured richness-based (R² = 0.135, p = 0.001), "
                "abundance-weighted (R² = 0.150, p = 0.001), and phylogenetic β-diversity (R² = 0.183, p = 0.001), "
                "but not functional β-diversity (R² = 0.005, p = 0.755; Fig. 2). Island additionally structured taxonomic "
                "β-diversity (PERMANOVA R² ≈ 0.13–0.15, p = 0.001) without corresponding effects on α-diversity. "
                "Compositional turnover therefore occurred primarily in taxonomic and phylogenetic space, with functional "
                "potential largely conserved across environments. Land-cover heterogeneity was a weak predictor of taxonomic "
                "β-diversity (R² = 0.027–0.045) and unrelated to functional turnover. NMDS ordinations mirrored these "
                "patterns (Fig. 2), with environmental groups partially separating for richness, neutral and phylogenetic "
                "distances—including clear mainland Tasmania separation—but showing substantial overlap in functional space.",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("Boxplots of possum gut microbiome"):
            replace_whole_paragraph(
                para,
                "Figure 2. Genome-resolved α- and β-diversity across broad environments. "
                "(a) Hill-number α-diversity (richness, neutral, phylogenetic and functional diversity) for 72 samples "
                "grouped by broad environment. (b) NMDS ordinations of pairwise β-diversity for taxonomic richness turnover, "
                "neutral turnover, phylogenetic turnover and functional turnover; points are coloured by broad environment "
                "and shaped by island (Tasmania vs mainland Australia).",
            )
            break

    # --- Results 3.3: HMSC scope intro (insert before devil paragraph) ---
    p = doc.paragraphs
    hmsc_scope = (
        "HMSC was fitted to 576 genomes in 55 Tasmania samples. Across predictors, most genomes were neutral and comparable "
        "minorities responded positively or negatively: devil density (130 positive, 105 negative, 341 neutral), temperature "
        "(139 positive, 134 negative, 303 neutral), and landscape diversity (133 positive, 109 negative, 334 neutral). "
        "A devil × temperature interaction was estimable (119 positive, 102 negative, 355 neutral), although geographic "
        "confounding limits inference about synergistic effects; 207 genomes showed congruent devil and temperature responses."
    )
    for i, para in enumerate(p):
        if para.text.strip().startswith("We then assessed whether microbiome features varied along the gradient"):
            insert_paragraph_before(para, hmsc_scope)
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("We then assessed whether microbiome features varied along the gradient"):
            replace_whole_paragraph(
                para,
                "We then assessed whether microbiome features varied along the gradient of Tasmanian devil density. "
                "In the final HMSC model, 130 genomes (23%) were positively associated with increasing devil density, "
                "105 (18%) were negatively associated, and 341 (59%) were neutral—a balanced pattern of partial association "
                "rather than community-wide enrichment. Devil-positive genomes were concentrated within Bacillota_A core "
                "families, especially Lachnospiraceae (n = 55) and Oscillospiraceae (n = 46). Devil-negative genomes included "
                "CAG-508 (n = 19), CAG-274 and CAG-465 among other lineages, indicating reweighting within the dominant "
                "phylum rather than wholesale turnover.",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("Using the HMSC-derived posterior estimates"):
            replace_substring(
                para,
                "Together, these patterns indicate that possums living in areas with high devil density harbour gut communities with a broader metabolic repertoire, including increased potential for amino acid and vitamin production, nitrogen recycling and the breakdown of complex carbohydrates and xenobiotics.",
                "Functional volcano contrasts identified 60 FDR-significant GIFT elements enriched among devil-positive genomes "
                "(Fig. 3), consistent with a broader metabolic repertoire under higher devil density, including amino acid and "
                "vitamin biosynthesis and polysaccharide, amino acid and xenobiotic degradation pathways.",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("HMSC revealed widespread but heterogeneous responses of individual MAGs"):
            replace_whole_paragraph(
                para,
                "HMSC revealed widespread but heterogeneous genome-level responses to host temperature: 139 genomes (24%) "
                "positive, 134 (23%) negative and 303 (53%) neutral. Warm-associated genomes were predominantly Bacillota_A "
                "(128 of 139; 92%), including Oscillospiraceae (n = 36) and Lachnospiraceae (n = 76); CAG-508 was not enriched "
                "among temperature-positive genomes. Cooler-associated genomes were more taxonomically diverse, including "
                "Cyanobacteriota (RUG14156), Thermoplasmatota (Methanomethylophilaceae) and Bacteroidota (Rikenellaceae), "
                "indicating taxonomic narrowing towards Firmicutes-dominated communities at warmer sites.",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("At the phylum level, Bacillota_A tended to show positive mean"):
            replace_whole_paragraph(
                para,
                "At the phylum level, Bacillota_A showed positive mean temperature coefficients, whereas Cyanobacteriota, "
                "Thermoplasmatota and several Bacteroidota and Pseudomonadota lineages were on average negatively associated "
                "with temperature. Among the strongest family-level responses, 36 Oscillospiraceae MAGs were temperature-positive "
                "compared with cooler-associated RUG14156, Methanomethylophilaceae and Rikenellaceae lineages.",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("Functional modelling identified a small set of metabolic pathways"):
            replace_whole_paragraph(
                para,
                "Functional modelling identified 61 FDR-significant GIFT elements differing between temperature-positive and "
                "temperature-negative genomes (Fig. 3). Temperature-positive genomes were enriched in biosynthetic and "
                "degradative traits, reversing the earlier expectation that cooler hosts would harbour broader metabolic "
                "potential at genome level. Spore formation (S0301) showed only marginal support (mean difference = +0.083, "
                "FDR = 0.045), so a broad transmission–survival signal was not strongly supported.",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("a)Phylum-level summary of HMSC devil density effects"):
            replace_whole_paragraph(
                para,
                "Figure 3. HMSC functional associations across four predictors (576 Tasmania genomes). Four-panel volcano plots "
                "show GIFT element differences between genomes classified as positively vs negatively associated with devil "
                "density, temperature, devil × temperature interaction, and landscape diversity; points are coloured by GIFT "
                "category and dashed lines indicate FDR = 0.05 and effect-size thresholds.",
            )
            break

    # --- Results 3.4 ---
    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("To exclude and get a clearer picture of the parameters influence"):
            replace_whole_paragraph(
                para,
                "The HMSC analyses above focused on Tasmania (55 samples, 576 genomes). Mean model fit was high across "
                "genomes (mean R² ≈ 0.77). Below we summarise taxonomic and functional patterns from phylum-level summaries "
                "and element-level volcano contrasts (Figs 3–4).",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("Devil density was associated with pronounced genome-level responses"):
            replace_substring(
                para,
                "Among the genomes with the strongest credible effects, responses were dominated by Bacillota and Bacillota_A, with a smaller subset of Bacteroidota contributing to the negative tail.",
                "Among genomes with credible effects, responses were dominated by Bacillota_A reweighting: comparable minorities "
                "were devil-positive (23%) or devil-negative (18%), with most genomes neutral (59%).",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("In contrast to devil density, temperature effects on individual genomes"):
            replace_whole_paragraph(
                para,
                "Temperature effects on individual genomes were comparable in magnitude to devil density (139 positive, 134 "
                "negative, 303 neutral). Functional contrasts identified 61 FDR-significant GIFT elements enriched among "
                "temperature-positive genomes, predominantly biosynthetic and degradative categories (Fig. 3). At the phylum "
                "level, Bacillota_A showed positive mean temperature coefficients whereas Thermoplasmatota and several "
                "Bacteroidota lineages tended negative. Overall, temperature reshaped which lineages dominated within a shared "
                "functional backdrop rather than reorganising community-wide functional potential.",
            )
            break

    # --- Discussion ---
    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("The one finding of this study is a marked decoupling"):
            replace_substring(
                para,
                "Although community composition and phylogenetic breadth shifted significantly across habitats and  geographic regions, aggregate functional potential remained homeostatic.",
                "Although community composition and phylogenetic β-diversity shifted significantly across habitats, islands and geographic regions, aggregate functional α- and β-diversity remained comparatively homeostatic.",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("Predation pressure from the Tasmanian devil"):
            replace_substring(
                para,
                "High-risk landscapes correlated with enrichment of biosynthetic pathways for vitamins, amino acids, and essential cofactors.",
                "Higher devil density was associated with genome-specific—not community-wide—enrichment of biosynthetic pathways for vitamins, amino acids, and essential cofactors among devil-positive MAGs (130 of 576; 23%), balanced by a comparable devil-negative subset (18%).",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("Ambient temperature was associated with a distinct trade-off"):
            replace_whole_paragraph(
                para,
                "Ambient temperature was associated with taxonomic narrowing towards Bacillota_A-dominated warm-associated "
                "genome sets (92% of temperature-positive MAGs), alongside enrichment of biosynthetic and degradative GIFT "
                "elements among temperature-positive genomes. Cooler-associated genomes were taxonomically more diverse "
                "(Cyanobacteriota, Thermoplasmatota, Bacteroidota). A broad spore-formation signal was weak (one marginal "
                "S-element), so transmission–survival (H3) remains plausible but is not strongly supported here.",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("The independent functional shifts associated with predation and thermal stress"):
            replace_whole_paragraph(
                para,
                "Devil density, temperature and landscape diversity represented parallel, partially independent axes of "
                "genome-level microbiome variation. A fitted devil × temperature interaction (119 positive, 102 negative "
                "genomes) and 207 congruent devil/temperature responses suggest overlapping stress contexts, but geographic "
                "confounding (devils only in Tasmania) limits inference about synergistic effects. Conservation management "
                "should therefore treat these gradients as associational filters on lineage composition rather than as validated "
                "causal levers for intervention.",
            )
            break

    p = doc.paragraphs
    for para in p:
        if para.text.strip().startswith("Although this study employs genome-resolved metagenomics"):
            replace_substring(
                para,
                "Although this study employs genome-resolved metagenomics to provide a high-resolution view of microbial functional potential, climate layers and predator density models represent coarse proxies for individual host experience.",
                "Although this study employs genome-resolved metagenomics to provide a high-resolution view of microbial functional potential, climate layers and predator density models represent coarse proxies for individual host experience. HMSC was restricted to 55 Tasmania samples (576 genomes) whereas α/β-diversity analyses used all 72 analysed samples (55 Tasmania, 17 mainland Australia).",
            )
            break

    doc.save(str(OUT))
    print(f"Backup: {BACKUP}")
    print(f"Revised manuscript saved: {OUT}")
    print("All new/revised text is coloured blue (RGB 0, 112, 192).")


if __name__ == "__main__":
    main()
