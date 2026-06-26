#!/usr/bin/env python3
"""Revise Discussion for devil density × temperature / physiological trap framing.

Reads:  Manuscript/4 | Discussion.docx
Backup: Manuscript/4 | Discussion_pre_physiological_trap.docx
Writes: Manuscript/4 | Discussion.docx

All new or changed wording is coloured blue (RGB 0, 112, 192).
"""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "Manuscript" / "4 | Discussion.docx"
BACKUP = ROOT / "Manuscript" / "4 | Discussion_pre_physiological_trap.docx"
OUT = SRC

BLUE = RGBColor(0, 112, 192)


def replace_whole_paragraph(paragraph, text: str, color: RGBColor = BLUE) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.color.rgb = color


def replace_substring(paragraph, old: str, new: str, color: RGBColor = BLUE) -> bool:
    text = paragraph.text
    if old not in text:
        return False
    before, _, after = text.partition(old)
    paragraph.clear()
    if before:
        paragraph.add_run(before)
    run_new = paragraph.add_run(new)
    run_new.font.color.rgb = color
    if after:
        paragraph.add_run(after)
    return True


def insert_paragraph_before(paragraph, text: str, color: RGBColor = BLUE):
    parent = paragraph._element.getparent()
    new_p = deepcopy(paragraph._element)
    parent.insert(parent.index(paragraph._element), new_p)
    from docx.text.paragraph import Paragraph

    new_para = Paragraph(new_p, paragraph._parent)
    replace_whole_paragraph(new_para, text, color=color)
    return new_para


def find_paragraph(doc, predicate):
    for para in doc.paragraphs:
        if predicate(para.text.strip()):
            return para
    return None


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Missing discussion document: {SRC}")

    shutil.copy2(SRC, BACKUP)
    doc = Document(str(SRC))

    # --- Remove thesis page-break artefacts (blue fixes) ---
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("73 captures")),
        "73 captures",
        "This catalogue captures",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("74 as a generalist")),
        "74 as a generalist",
        "As a generalist",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("75 structure.")),
        "75 structure. A substantial proportion of genomes exhibited positive associations with increasing devil density.",
        (
            "structure. Among 576 MAGs modelled across 55 Tasmania samples, comparable minorities of genomes were "
            "devil-positive (130; 23%) or devil-negative (105; 18%), with most neutral (341; 59%)."
        ),
    )
    replace_substring(
        find_paragraph(doc, lambda t: "76 Q3 and allows" in t),
        "76 Q3 and allows",
        "Q3 and allows",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("77 flux.")),
        "77 flux.",
        "",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("78 by high posterior")),
        "78 by high posterior",
        "By high posterior",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("79 or longitudinal")),
        "79 or longitudinal",
        "Or longitudinal",
    )

    # --- §4.2: island/alpha wording ---
    replace_substring(
        find_paragraph(
            doc,
            lambda t: t.startswith(
                "A central finding of this study, and the answer to Q1"
            ),
        ),
        (
            "Habitat categories and island of origin (Tasmania versus mainland Australia) explained "
            "variation in taxonomic alpha and beta diversity, with ordination analyses revealing "
            "distinct clustering of samples by environment in both neutral and phylogenetic spaces."
        ),
        (
            "Broad environment categories explained variation in taxonomic alpha diversity, whereas "
            "island of origin (Tasmania versus mainland Australia) and habitat jointly structured "
            "taxonomic beta diversity, with ordination revealing distinct clustering in neutral and "
            "phylogenetic spaces."
        ),
    )

    # --- §4.2 bridge to stress gradients ---
    replace_substring(
        find_paragraph(
            doc,
            lambda t: t.endswith(
                "I now turn to a specific ecological axis: predation pressure and its downstream influence on microbiome structure."
            ),
        ),
        (
            "Having established this general pattern of taxonomic flexibility against functional conservatism, "
            "I now turn to a specific ecological axis: predation pressure and its downstream influence on microbiome structure."
        ),
        (
            "Crucially, this community-level homeostasis does not preclude genome-specific compensation along "
            "environmental stress gradients: the holobiont may maintain a stable functional ceiling while "
            "reweighting which lineages carry biosynthetic and detoxification load. I therefore turn next to "
            "predation pressure and temperature as complementary axes shaping microbiome structure."
        ),
    )

    # --- §4.3 intro: tighten duplicate species name ---
    replace_substring(
        find_paragraph(
            doc,
            lambda t: t.startswith(
                "To identify whether predation risk drives metagenomic shifts (Q2)"
            ),
        ),
        (
            "Predation pressure exerted by the Tasmanian devil (Sarcophilus harrisii) was hypothesized to influence "
            "the possum holobiont indirectly through behavioural modifications, such as altered habitat utilization, "
            "constrained foraging ranges, and subsequent dietary or physiological shifts. The HMSC results provide "
            "robust evidence that this ecological gradient is reflected in the microbiome"
        ),
        (
            "Predation pressure was hypothesized to influence the possum holobiont indirectly through behavioural "
            "modifications—altered habitat use, constrained foraging ranges, and subsequent dietary or physiological "
            "shifts. HMSC provides robust evidence that this gradient is reflected in genome-level microbiome"
        ),
    )

    # --- §4.3 main devil paragraph (remainder after opening fix) ---
    p_devil = find_paragraph(doc, lambda t: "devil-positive (130; 23%)" in t or "devil-positive" in t and "Bacillota_A families" in t)
    if p_devil is None:
        p_devil = find_paragraph(doc, lambda t: "Bacillota_A families that dominate" in t)
    if p_devil is not None:
        replace_substring(
            p_devil,
            (
                "These associations were concentrated within the Bacillota_A families that dominate the core community, "
                "while a smaller subset of genomes displayed negative associations. Functionally, genomes associated "
                "with higher devil densities were enriched in both biosynthetic and degradative capacities."
            ),
            (
                "Responses were concentrated within Bacillota_A core families—notably Lachnospiraceae and "
                "Oscillospiraceae—indicating reweighting within the dominant guild rather than wholesale turnover. "
                "Among devil-positive genomes, 60 GIFT elements differed significantly from devil-negative genomes "
                "(FDR < 0.05), with enrichment of both biosynthetic and degradative capacities."
            ),
        )

    # --- §4.4 heading ---
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("4.4 | Temperature")),
        "4.4 | Temperature, Firmicutes, and Cold-Associated Functional Breadth",
        "4.4 | Temperature, Firmicutes, and Genome-Resolved Thermal Filtering",
    )

    # --- §4.4 temperature body ---
    replace_whole_paragraph(
        find_paragraph(doc, lambda t: "Q3 and allows assessment of the transmission-survival" in t),
        (
            "Q3 and allows assessment of the transmission-survival hypothesis (H3). Unlike broad habitat categories, "
            "temperature aligned with distinct genome-level contrasts. Temperature-positive genomes were predominantly "
            "Bacillota_A (128 of 139; 92%), especially Oscillospiraceae, whereas cooler-associated genomes were "
            "taxonomically more diverse, including Cyanobacteriota, Thermoplasmatota, and Bacteroidota lineages. "
            "Comparable minorities were temperature-positive (139; 24%) or temperature-negative (134; 23%), mirroring "
            "the magnitude of devil-associated filtering. Functional volcano contrasts identified 61 FDR-significant "
            "GIFT elements enriched among temperature-positive genomes, predominantly biosynthetic and degradative "
            "categories—rather than a broader metabolic portfolio in cool-associated genomes. The shift toward "
            "Bacillota_A at higher temperatures is consistent with diet-induced thermogenesis logic: as possums reduce "
            "protein intake in the heat to limit internal overheating (Beale et al., 2023), the microbiome may pivot "
            "toward Firmicutes lineages efficient at carbohydrate fermentation (Beale et al., 2022; Youngentob et al., "
            "2021). Directional genome responses may also buffer host physiological stress (Fontaine et al., 2023; "
            "Koziol et al., 2023). A broad spore-formation signal was weak (one marginal S-element), so "
            "transmission-survival (H3) remains plausible but is not strongly supported here. These patterns emphasize "
            "that inference rests on genomic potential rather than measured gene expression or metabolite flux."
        ),
    )

    # --- §4.4 bridge to trap ---
    replace_whole_paragraph(
        find_paragraph(
            doc,
            lambda t: t.startswith("These independent associations with devil density and temperature")
            or t.startswith(" flux. These independent associations"),
        ),
        (
            "These independent associations with devil density and temperature raise a broader concern: in landscapes "
            "where both pressures coincide, possums may face compounding constraints that the microbiome alone cannot "
            "fully buffer. Functional redundancy at community scale may therefore mask genome-specific compensation "
            "that becomes critical when predator recovery and warming intensify together."
        ),
    )

    # --- §4.4.1 heading ---
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("4.4.1 | Predation")),
        "4.4.1 | Predation x Temperature: Evidence Consistent with a Physiological Trap",
        "4.4.1 | Devil Density × Temperature: Evidence Consistent with a Physiological Trap",
    )

    # --- §4.4.1 body: replace + insert two supporting paragraphs ---
    p_trap = find_paragraph(
        doc,
        lambda t: t.startswith(
            "To examine whether the synergistic effects of warming and predator recovery"
        ),
    )
    if p_trap is None:
        raise SystemExit("Could not find §4.4.1 paragraph.")

    replace_whole_paragraph(
        p_trap,
        (
            "To examine whether warming and predator recovery jointly constrain possum populations—making microbiome "
            "mediation indispensable (Q4)—I evaluated devil density, temperature, and their interaction in the final "
            "HMSC model. Chapter 1 proposed that a physiological trap could emerge when predators restrict behavioural "
            "thermoregulation and foraging options while heat stress concurrently alters dietary requirements and "
            "detoxification demand (Beale et al., 2022; Beale et al., 2023; Youngentob et al., 2021). Both predictors "
            "showed strong independent genome-level associations, and their directionality is broadly consistent with "
            "this compensatory logic."
        ),
    )

    p_sec45 = find_paragraph(doc, lambda t: t.startswith("4.5 | HMSC"))
    if p_sec45 is None:
        raise SystemExit("Could not find §4.5 heading.")

    for text in (
        (
            "Strikingly, among the 207 genomes with non-neutral responses to both devil density and temperature, all were "
            "directionally congruent: 112 were devil-positive and temperature-positive, and 95 were devil-negative and "
            "temperature-negative, with no discordant devil-positive/temperature-negative or devil-negative/"
            "temperature-positive pairs. This pattern indicates parallel filtering of a shared stress-responsive core "
            "within Bacillota_A rather than independent turnover along orthogonal axes."
        ),
        (
            "The fitted devil × temperature interaction identified a further layer of non-additivity: 119 genomes (21%) "
            "showed a positive interaction coefficient and 102 (18%) a negative interaction. Interaction-positive "
            "genomes were concentrated among devil-negative lineages (102 of 119), enriched in nitrogen-compound "
            "degradation, amino-acid degradation, and vitamin biosynthesis among 16 FDR-significant GIFT contrasts. "
            "Conversely, 96 of 102 interaction-negative genomes were devil-positive and temperature-positive—lineages "
            "recruited under both main-effect gradients whose joint responses fell short of additive expectation. "
            "Together, congruent main effects and antagonistic interaction coefficients are compatible with compensatory "
            "reweighting that may approach limits when both stressors intensify together."
        ),
        (
            "Interpretation must remain cautious. Devil density and temperature are spatially structured across Tasmania "
            "(r ≈ −0.19 across 55 HMSC samples), and devils occur only on the island; the evidence therefore supports "
            "H4 as directional consistency rather than confirmation of synergistic effects. Targeted sampling across "
            "warm/dry versus cool/wet regions with contrasting devil densities—and paired diet or stress biomarkers—"
            "would be needed to test the trap hypothesis causally."
        ),
    ):
        insert_paragraph_before(p_sec45, text)

    # --- §4.5: remove DAMR claim ---
    replace_whole_paragraph(
        find_paragraph(
            doc,
            lambda t: t.startswith("In addition to advanced modelling, technical validation")
        ),
        (
            "Technical validation was supported through the SingleM Microbial Fraction (SMF) tool, which allowed "
            "estimation of microbial abundance for novel lineages missing from reference databases (Eisenhofer et al., "
            "2024). This supports the view that functional inferences rest on a genome-resolved metagenomic landscape "
            "(Odriozola et al., 2024)."
        ),
    )

    # --- §4.6: expand limitations ---
    replace_substring(
        find_paragraph(
            doc,
            lambda t: t.startswith("Several limitations must be considered"),
        ),
        (
            "Second, the sampling design was imbalanced across regions, with the warm end of the temperature gradient "
            "represented by fewer mainland samples compared to the extensive Tasmanian dataset. This reduces statistical "
            "power for climate-related inferences and interaction testing."
        ),
        (
            "Second, the sampling design was imbalanced across regions, with the warm end of the temperature gradient "
            "represented by fewer mainland samples compared to the extensive Tasmanian dataset, and HMSC restricted to "
            "55 Tasmania samples (576 genomes) whereas α/β-diversity analyses used all 72 analysed samples. Devil density "
            "and temperature are additionally partially confounded geographically, limiting robust inference on their "
            "interaction."
        ),
    )

    # --- §4.7 intro ---
    replace_whole_paragraph(
        find_paragraph(
            doc,
            lambda t: t.startswith("While the core functional stability supports the idea of functional redundancy")
        ),
        (
            "While core functional stability supports functional redundancy, targeted genome-level shifts along devil "
            "and temperature gradients—biosynthesis, degradation, and detoxification—are consistent with compensatory "
            "roles that may become limiting when both stressors coincide."
        ),
    )

    # --- §4.7.1 conservation ---
    replace_whole_paragraph(
        find_paragraph(
            doc,
            lambda t: t.startswith("These findings support the utility of the holobiont concept")
        ),
        (
            "These findings support the holobiont concept for anticipating wildlife responses to environmental change. "
            "If predator recovery and climatic warming intensify simultaneously, persistence may depend on whether "
            "recipient populations carry microbiome configurations capable of biosynthetic and detoxification "
            "compensation—a hologenomic extension of translocation risk assessment (Molloy et al., 2016). Experimental "
            "validation (e.g., diet metabarcoding, stress biomarkers, or microbiota manipulations) should precede any "
            "management interventions such as probiotics or bioaugmentation."
        ),
    )

    # --- §4.7.4 take-home ---
    replace_whole_paragraph(
        find_paragraph(
            doc,
            lambda t: t.startswith("Despite limitations, this thesis reaches a coherent conclusion")
        ),
        (
            "Despite limitations, this thesis reaches a coherent conclusion: T. vulpecula hosts a gut microbiome that is "
            "taxonomically responsive yet functionally robust at community scale. Environmental gradients restructure "
            "lineage composition—especially within Bacillota_A—while genome-resolved HMSC reveals parallel devil- and "
            "temperature-associated compensation that is directionally consistent with, but does not confirm, a dual "
            "physiological trap when predation and warming coincide."
        ),
    )

    # --- Fix sentence breaks introduced by thesis page artefacts ---
    replace_substring(
        find_paragraph(doc, lambda t: t.endswith("genome-level microbiome")),
        "genome-level microbiome",
        "genome-level microbiome structure.",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("structure. Among 576 MAGs")),
        "structure. Among 576 MAGs",
        "Among 576 MAGs",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.endswith("genome responses, indicated")),
        "genome responses, indicated",
        "genome responses—indicated",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("By high posterior support")),
        "By high posterior support",
        "by high posterior support",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.endswith("experimental manipulations")),
        "experimental manipulations",
        "experimental manipulations or longitudinal studies would be required to confirm causal pathways.",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("Or longitudinal studies")),
        "Or longitudinal studies would be required to confirm causal pathways. ",
        "",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("80 depletion")),
        "80 depletion, to establish",
        "depletion—to establish",
    )
    replace_substring(
        find_paragraph(
            doc,
            lambda t: "microbiota transplants or antibiotic" in t
            and not t.endswith("depletion—to establish"),
        ),
        "microbiota transplants or antibiotic",
        "microbiota transplants or antibiotic depletion—to establish causal roles for "
        "microbiome-mediated environmental buffering (Kolodny & Schulenburg, 2020).",
    )
    p_depl = find_paragraph(doc, lambda t: t.startswith("depletion—to establish"))
    if p_depl is not None and p_depl.text.strip().startswith("depletion—to establish"):
        replace_whole_paragraph(p_depl, "")

    # --- §4.5: repair split phylogenetic-signal sentence ---
    p_hmsc = find_paragraph(
        doc,
        lambda t: t.startswith("While ecological signals were detectable")
        and t.endswith("indicated"),
    )
    p_hmsc_cont = find_paragraph(
        doc, lambda t: t.startswith("by high posterior support for phylogenetically")
    )
    if p_hmsc is not None and p_hmsc_cont is not None:
        replace_whole_paragraph(
            p_hmsc,
            (
                "While ecological signals were detectable using standard ordination and PERMANOVA "
                "techniques, the HMSC framework provided a significantly deeper, genome-resolved "
                "perspective. This approach allowed simultaneous evaluation of community covariance "
                "with multiple predictors while accounting for hierarchical study design and shared "
                "evolutionary history among genomes. A strong phylogenetic signal in genome responses, "
                "indicated by high posterior support for phylogenetically structured effects, suggests "
                "that environmental sensitivities are not randomly distributed across the microbial tree "
                "of life but are conserved within specific clades. This has important interpretive "
                "implications: shifts along ecological gradients are not merely changes in species "
                "richness, but structured reweighting of phylogenetically coherent groups possessing "
                "correlated trait repertoires. Joint modelling thus complements traditional beta-diversity "
                "metrics by allowing mechanistic hypotheses to be evaluated at the level of individual "
                "genomes and trait bundles."
            ),
        )
        replace_whole_paragraph(p_hmsc_cont, "")

    # --- §4.4: fix split temperature intro ---
    replace_substring(
        find_paragraph(
            doc,
            lambda t: t.startswith("Devil density captures one axis") and t.endswith("addresses"),
        ),
        "along this abiotic axis addresses",
        "along this abiotic axis addresses Q3.",
    )
    replace_substring(
        find_paragraph(doc, lambda t: t.startswith("Q3 and allows assessment")),
        "Q3 and allows assessment of the transmission-survival hypothesis (H3). Unlike",
        "Assessment of the transmission-survival hypothesis (H3) follows. Unlike",
    )

    doc.save(str(OUT))
    print(f"Backup: {BACKUP}")
    print(f"Revised discussion saved: {OUT}")
    print("All new/changed text is coloured blue (RGB 0, 112, 192).")


if __name__ == "__main__":
    main()
