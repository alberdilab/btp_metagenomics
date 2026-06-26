#!/usr/bin/env python3
"""Third pass: typos, section numbering, and figure-caption standardisation.

New/changed text coloured teal (RGB 0, 128, 128).
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
MANUSCRIPT = ROOT / "Manuscript" / "btp_manuscript.docx"
BACKUP = ROOT / "Manuscript" / "btp_manuscript_pre_polish_pass3.docx"

TEAL = RGBColor(0, 128, 128)


def find_paragraph(doc: Document, predicate):
    for para in doc.paragraphs:
        if predicate(para.text.strip()):
            return para
    raise ValueError("Paragraph not found")


def write_teal(para: Paragraph, text: str, *, style: str | None = None) -> None:
    para.clear()
    if style:
        para.style = style
    if text:
        run = para.add_run(text)
        run.font.color.rgb = TEAL


def insert_heading_before(anchor: Paragraph, text: str, *, style: str = "Heading 2") -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._element.getparent().insert(anchor._element.getparent().index(anchor._element), new_p)
    para = Paragraph(new_p, anchor._parent)
    write_teal(para, text, style=style)
    return para


# (predicate, text, optional_style_override)
BODY_REPLACEMENTS: list[tuple[object, str, str | None]] = [
    (
        lambda t: t.startswith("3.3 | Metabolic capacitys"),
        "3.3 | Genome-resolved HMSC associations and functional context",
        "Heading 2",
    ),
    (
        lambda t: t.startswith("The assembly and binning of 352.3 Gb"),
        "The assembly and binning of 352.3 Gb (4.89 ± 1.41 Gb/sample) of metagenomic data yielded "
        "1,817 non-redundant MAGs spanning 13 bacterial and 2 archaeal phyla (Fig. 1). Read "
        "partitioning revealed that 70.9 ± 8.6% of reads mapped to microbial MAGs, 19.3 ± 6.0% "
        "remained unmapped, 7.3 ± 3.4% were classified as low-quality, and 2.5 ± 8.4% mapped to "
        "the host genome. Average CheckM quality of MAGs was 79.4 ± 16.8% completeness and "
        "2.79 ± 2.58% contamination. One sample (EHI01340) was excluded for insufficient MAG reads.",
        None,
    ),
    (
        lambda t: t.startswith("Figure 1: Catalog of brushtail-associated microbial genoms"),
        "Fig. 1. Catalogue of brushtail-associated microbial genomes. (A) Reference "
        "metagenome-assembled genome (MAG) catalogue from 72 samples. Outer rings indicate MAG "
        "detection in different environmental clusters and genome size.",
        None,
    ),
    (
        lambda t: t.startswith("Figure 2: Study context and taxonomic composition"),
        "Fig. 2. Study context and taxonomic composition. (A) Map of sample locations in Australia "
        "and Tasmania. Left, full sample set; right, Tasmanian samples; dot size indicates sample "
        "number clustered into four environments; green ring intensity indicates devil density. "
        "(B) MAG-level taxonomic composition tile plot, coloured by phylum.",
        None,
    ),
    (
        lambda t: t.startswith("Fig3: α and β-diversity"),
        "Fig. 3. α and β-diversity across broad environments. (A) MAG richness, neutral diversity, "
        "phylogenetic diversity and functional diversity, with individual samples as points and "
        "boxes summarising habitat-level distributions. (B) NMDS ordination of gut microbiome "
        "β-diversity; points represent individuals, coloured by broad environmental context and "
        "shaped by region.",
        None,
    ),
    (
        lambda t: t.startswith("To place this genome-level HMSC associations"),
        "To place these genome-level HMSC associations in phylogenetic and functional context, we "
        "visualised a stratified subset of responsive MAGs (n = 120) on a pruned phylogeny with a "
        "linked annotation heatmap (Fig. 4). Most responsive lineages fell within Bacillota_A and "
        "related Firmicutes clades, with smaller contributions from Bacteroidota, Cyanobacteriota "
        "and other phyla. Adjacent HMSC strips display scaled posterior beta coefficients, showing "
        "ecological associations are phylogenetically structured but not phylum-uniform. Closely "
        "related genomes often share similar response signs, yet sister lineages can diverge in "
        "direction. At the function-family level, biosynthetic and degradative categories form the "
        "dominant axes of variation among genomes.",
        None,
    ),
    (
        lambda t: t.startswith("fig4) Phylogenetic tree and heatmap"),
        "Fig. 4. Phylogenetic tree and heatmap of reconstructed genomes, highlighting significant "
        "association with devil density, temperature, landscape diversity or devil × temperature "
        "interaction, and capacity to perform metabolic functions within each GIFT group. "
        "Fine-grained visualisations of metabolic capacities are provided in the Supplementary "
        "Information.",
        None,
    ),
    (
        lambda t: "Devil" in t and "associated shifts" in t and "possum" in t and not t.startswith("3.4"),
        "3.4.1 | Devil-associated shifts in the microbiome",
        "Heading 4",
    ),
    (
        lambda t: t.startswith("The final HMSC model showed a balanced pattern"),
        "The final HMSC model showed a balanced pattern of partial association rather than "
        "community-wide enrichment. Devil-positive genomes were concentrated within Bacillota_A core "
        "families, especially Lachnospiraceae (n = 55) and Oscillospiraceae (n = 46). "
        "Devil-negative genomes included CAG-508 (n = 19), CAG-274 and CAG-465 among other lineages, "
        "indicating reweighting within the dominant phylum rather than wholesale turnover.",
        None,
    ),
    (
        lambda t: "Temperature" in t and "associated shifts" in t and "possum" in t and not t.startswith("3.4"),
        "3.4.2 | Temperature-associated shifts in the microbiome",
        "Heading 4",
    ),
    (
        lambda t: t.startswith("HMSC revealed widespread but heterogeneous responses to environmental temperature"),
        "HMSC revealed widespread but heterogeneous responses to environmental temperature. "
        "Warm-associated MAGs belonged predominantly to Bacillota_A (128 of 139; 92%), in particular "
        "Lachnospiraceae (n = 76) and Oscillospiraceae (n = 36). Cooler-associated genomes were more "
        "taxonomically diverse, concentrated in Cyanobacteriota (RUG14156 and Gastranaerophilaceae), "
        "Thermoplasmatota (Methanomethylophilaceae), Bacteroidota (Rikenellaceae) and Pseudomonadota "
        "(Burkholderiaceae_A), indicating taxonomic narrowing towards Firmicutes-dominated communities "
        "at warmer sites.",
        None,
    ),
    (
        lambda t: t.startswith("3.4.3 Interaction-associated"),
        "3.4.3 | Interaction-associated shifts in the microbiome",
        "Heading 4",
    ),
    (
        lambda t: t.startswith("3.4.4 Diversity-associated"),
        "3.4.4 | Diversity-associated shifts in the microbiome",
        "Heading 4",
    ),
]

SECTION_34_HEADING = "3.4 | Functional shifts associated with environmental predictors"


def main() -> None:
    shutil.copy2(MANUSCRIPT, BACKUP)
    doc = Document(str(MANUSCRIPT))

    devil_heading = find_paragraph(
        doc,
        lambda t: "Devil" in t and "associated shifts" in t and "possum" in t and not t.startswith("3.4"),
    )
    if not any(p.text.strip() == SECTION_34_HEADING for p in doc.paragraphs):
        insert_heading_before(devil_heading, SECTION_34_HEADING, style="Heading 2")

    n = 0
    for pred, text, style in BODY_REPLACEMENTS:
        para = find_paragraph(doc, pred)
        write_teal(para, text, style=style)
        n += 1

    doc.save(str(MANUSCRIPT))
    print(f"Backup: {BACKUP}")
    print(f"Polished {n} blocks + section 3.4 heading (teal).")


if __name__ == "__main__":
    main()
